
import openpyxl
from pathlib import Path
import sys

# Force UTF-8 for output
sys.stdout.reconfigure(encoding='utf-8')

base_dir = Path("c:/laragon/www/Rekap Realisasi PAD WIL JAWA 2021-2025")
input_file = base_dir / "assets" / "datarekaprealisasi.xlsx"

try:
    wb = openpyxl.load_workbook(input_file, read_only=True)
    if 'LRA (sort) (2)' in wb.sheetnames:
        sheet = wb['LRA (sort) (2)']
        print(f"Sheet: {sheet.title}")
        for i, row in enumerate(sheet.iter_rows(min_row=5, max_row=8, values_only=True)):
            clean_row = [str(cell) for cell in row]
            print(f"Row {i+5}: {clean_row}")
    else:
        print("Sheet not found")
except Exception as e:
    print(f"Error: {e}")

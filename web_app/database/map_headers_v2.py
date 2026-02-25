import openpyxl

path = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets\REALISASI PENDAPATAN TA 2025.xlsx'
wb = openpyxl.load_workbook(path, read_only=True)
ws = wb[wb.sheetnames[-1]]

print(f"Reading rows 10-12 from {path}")
for i, row in enumerate(ws.iter_rows(min_row=10, max_row=12, values_only=True), 10):
    print(f"--- Row {i} ---")
    for j, val in enumerate(row):
        if val:
            v_str = str(val).strip().replace('\n', ' ')
            if len(v_str) > 20: v_str = v_str[:20] + '...'
            print(f"{j}: {v_str}")
wb.close()

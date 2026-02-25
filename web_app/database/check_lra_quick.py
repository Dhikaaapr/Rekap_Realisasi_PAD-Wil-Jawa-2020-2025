import requests
import openpyxl

SUPABASE_URL = 'https://xerkytrweahuniqrnabi.supabase.co'
SUPABASE_KEY = 'sb_publishable_IhhX55NqvJ7bUIoXG_xxCw_BM66VfG7'

headers = {
    'apikey': SUPABASE_KEY,
    'Authorization': 'Bearer ' + SUPABASE_KEY,
    'Content-Type': 'application/json',
    'Prefer': 'count=exact'
}

out = []
for year in [2021, 2022, 2023, 2024, 2025]:
    all_details = []
    offset = 0
    while True:
        resp = requests.get(
            SUPABASE_URL + '/rest/v1/detail_pad_data?select=daerah,kategori_kode&tahun=eq.' + str(year) + '&order=daerah&offset=' + str(offset) + '&limit=1000',
            headers=headers
        )
        batch = resp.json()
        if not batch:
            break
        all_details.extend(batch)
        if len(batch) < 1000:
            break
        offset += 1000
    
    regions = {}
    for d in all_details:
        r = d['daerah']
        regions[r] = regions.get(r, 0) + 1
    
    granular = sum(1 for c in regions.values() if c > 4)
    basic = sum(1 for c in regions.values() if c <= 4)
    
    out.append(f"{year}: {len(all_details)} rows, {len(regions)} regions, {granular} granular, {basic} basic")

# Check sheets
kabkota_path = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets\KABKOTA.xlsx'
wb = openpyxl.load_workbook(kabkota_path, read_only=True)
out.append(f"KABKOTA sheets: {wb.sheetnames}")
wb.close()

prov_path = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets\PROVINSI.xlsx'
wb2 = openpyxl.load_workbook(prov_path, read_only=True)
out.append(f"PROVINSI sheets: {wb2.sheetnames}")
wb2.close()

# Write as plain ASCII
with open('lra_result.md', 'w', encoding='ascii', errors='replace') as f:
    for line in out:
        f.write(line + '\n')

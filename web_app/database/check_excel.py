import openpyxl
import os

assets = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets'

# 1. REALISASI PENDAPATAN TA 2025 - last sheet
print("=" * 60)
print("REALISASI PENDAPATAN TA 2025")
print("=" * 60)
wb = openpyxl.load_workbook(os.path.join(assets, 'REALISASI PENDAPATAN TA 2025.xlsx'), read_only=True)
print(f'Sheets: {wb.sheetnames}')
ws = wb[wb.sheetnames[-1]]
print(f'Last sheet: {wb.sheetnames[-1]}, rows={ws.max_row}, cols={ws.max_column}')
for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, max_col=40, values_only=True), 1):
    non_empty = [(j, str(v)[:50]) for j, v in enumerate(row) if v is not None]
    if non_empty:
        print(f'R{i}: {non_empty}')
wb.close()

# 2. PROVINSI.xlsx - first sheet  
print("\n" + "=" * 60)
print("PROVINSI.xlsx")
print("=" * 60)
wb2 = openpyxl.load_workbook(os.path.join(assets, 'PROVINSI.xlsx'), read_only=True)
print(f'Sheets: {wb2.sheetnames}')
ws2 = wb2[wb2.sheetnames[0]]
print(f'Sheet: {wb2.sheetnames[0]}, rows={ws2.max_row}, cols={ws2.max_column}')
for i, row in enumerate(ws2.iter_rows(min_row=1, max_row=15, max_col=30, values_only=True), 1):
    non_empty = [(j, str(v)[:50]) for j, v in enumerate(row) if v is not None]
    if non_empty:
        print(f'R{i}: {non_empty}')
wb2.close()

# 3. KABKOTA.xlsx - first sheet
print("\n" + "=" * 60)
print("KABKOTA.xlsx")
print("=" * 60)
wb3 = openpyxl.load_workbook(os.path.join(assets, 'KABKOTA.xlsx'), read_only=True)
print(f'Sheets: {wb3.sheetnames}')
ws3 = wb3[wb3.sheetnames[0]]
print(f'Sheet: {wb3.sheetnames[0]}, rows={ws3.max_row}, cols={ws3.max_column}')
for i, row in enumerate(ws3.iter_rows(min_row=1, max_row=15, max_col=30, values_only=True), 1):
    non_empty = [(j, str(v)[:50]) for j, v in enumerate(row) if v is not None]
    if non_empty:
        print(f'R{i}: {non_empty}')
wb3.close()

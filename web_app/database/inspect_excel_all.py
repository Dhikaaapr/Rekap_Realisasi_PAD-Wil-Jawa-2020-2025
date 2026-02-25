import openpyxl

out = []

# Check REALISASI PENDAPATAN files
for fname in ['REALISASI PENDAPATAN TA 2023.xlsx', 'REALISASI PENDAPATAN TA 2024.xlsx', 'REALISASI PENDAPATAN TA 2025.xlsx']:
    fpath = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets' + '\\' + fname
    try:
        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        out.append(f"File: {fname}")
        out.append(f"Sheets: {wb.sheetnames}")
        for sn in wb.sheetnames[:2]:
            ws = wb[sn]
            row_count = 0
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), 1):
                clean = [str(c).strip().replace('\n', ' ')[:40] for c in row if c is not None]
                if clean:
                    out.append(f"  Row {i}: {clean[:8]}")
                row_count = i
            # Count total rows
            total = 0
            for _ in ws.iter_rows(values_only=True):
                total += 1
            out.append(f"  Total rows: {total}")
        wb.close()
    except Exception as e:
        out.append(f"Error with {fname}: {e}")
    out.append("")

# Check Database Rekap 2025
fpath = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets\Database Rekap 21 Des 2025 - TA 2025 1.xlsx'
try:
    wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
    out.append(f"File: Database Rekap 2025")
    out.append(f"Sheets: {wb.sheetnames}")
    for sn in wb.sheetnames:
        ws = wb[sn]
        total = sum(1 for _ in ws.iter_rows(values_only=True))
        out.append(f"  Sheet: {sn}, Total rows: {total}")
    wb.close()
except Exception as e:
    out.append(f"Error: {e}")

with open('excel_struct.py', 'w', encoding='ascii', errors='replace') as f:
    for line in out:
        f.write(line + '\n')
print("Done")

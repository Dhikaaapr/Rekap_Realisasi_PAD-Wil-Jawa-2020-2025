import openpyxl

path = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets\REALISASI PENDAPATAN TA 2025.xlsx'
print(f"Reading {path} (data_only=True)...")

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[-1]] # Last sheet
print(f"Sheet: {ws.title}")

for i, row in enumerate(ws.iter_rows(min_row=5, max_row=15, values_only=True), 5):
    # Print non-empty 
    vals = []
    for j, v in enumerate(row):
        if v:
            vals.append(f"{j}:{str(v).strip().replace('\n',' ')[:20]}")
    if vals:
        print(f"Row {i}: {vals}")

wb.close()

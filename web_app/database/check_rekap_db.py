import openpyxl
import sys

path = r'c:\laragon\www\Rekap Realisasi PAD WIL JAWA 2021-2025\assets\Database Rekap 21 Des 2025 - TA 2025 1.xlsx'
print(f"Checking {path}")

try:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    
    # Print first 10 rows
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        clean_row = [str(cell).strip().replace('\n', ' ')[:30] for cell in row if cell]
        print(f"Row {i}: {clean_row}")
        
    wb.close()
except Exception as e:
    print(f"Error: {e}")
